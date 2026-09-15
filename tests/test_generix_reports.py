import csv
import json
from pathlib import Path

from openpyxl import load_workbook

from integration_app.generix.reports import export_header_report


def test_export_header_report_writes_csv_and_json(tmp_path: Path):
    root = tmp_path / "Generix" / "Bat" / "storage" / "cpip_1"
    received = root / "received" / "headers"
    sent = root / "sent" / "headers"
    sent_data = root / "sent" / "data"
    received.mkdir(parents=True)
    sent.mkdir(parents=True)
    sent_data.mkdir(parents=True)
    edi_path = sent_data / "tx.txt"
    edi_path.write_text(
        """CAB220 202600552                          PT5106785059125042                 PT5066985992794093                 EntregaFarm- Logística FarmacêuticaPT5106785059125042                 ALCOTRANS (M)
DET0000015852355                                                               DIOFLAV 1000 MG 30 COMP
DET0000025852363                                                               DIOFLAV 1000 MG 60 COMP
TOT00000000002
""",
        encoding="utf-8",
    )
    (received / "rx.hdr").write_text(
        """date=Fri, 04 Sep 2026 09:10:11 +0100
from=farmacia@netixone.com.pt
to=influe@netixone.com.pt
subject=ord5139292776.txt
unique-id=NXC-RX
[log]
2026-09-04 09:10:12 message received.
""",
        encoding="utf-8",
    )
    (sent / "tx.hdr").write_text(
        f"""body-path={edi_path}
date=Sat, 12 Sep 2026 10:53:17 +0100
from=influe@netixone.com.pt
to=fch@netixone.com.pt
subject=Encomenda_20260912_105317.txt
unique-id=NXC-TX
disposition=automatic-action/MDN-sent-automatically; processed
[log]
2026-09-12 10:53:19 the message was PROCESSED
""",
        encoding="utf-8",
    )

    written = export_header_report(root, tmp_path / "reports", run_id="generix-test")

    assert written == [
        tmp_path / "reports" / "generix-test.csv",
        tmp_path / "reports" / "generix-test.json",
        tmp_path / "reports" / "generix-test-exceptions.csv",
        tmp_path / "reports" / "generix-test-exceptions.json",
        tmp_path / "reports" / "generix-test.xlsx",
    ]
    with written[0].open("r", newline="", encoding="utf-8") as handle:
        csv_rows = list(csv.DictReader(handle))
    json_rows = json.loads(written[1].read_text(encoding="utf-8"))
    assert [row["flow_type"] for row in csv_rows] == ["received", "sent"]
    assert json_rows[1]["unique_id"] == "NXC-TX"
    assert json_rows[1]["processed"] is True
    assert json_rows[1]["edi_path"] == str(edi_path)
    assert json_rows[1]["edi_detail_count"] == 2
    assert json_rows[1]["edi_has_total"] is True
    assert json_rows[1]["edi_origin_name"] == "EntregaFarm- Logística Farmacêutica"
    assert "PT5106785059125042" in json_rows[1]["edi_gln_codes"]
    assert json_rows[1]["exception_level"] == "ok"
    assert json_rows[1]["exception_reason"] == ""


def test_export_header_report_flags_exceptions(tmp_path: Path):
    root = tmp_path / "storage" / "cpip_1"
    sent = root / "sent" / "headers"
    sent_data = root / "sent" / "data"
    received = root / "received" / "headers"
    sent.mkdir(parents=True)
    sent_data.mkdir(parents=True)
    received.mkdir(parents=True)
    (sent_data / "missing-total.txt").write_text(
        "CAB220 202600552                          EntregaFarm\nDET000001\n",
        encoding="utf-8",
    )
    (sent / "missing-total.hdr").write_text(
        """unique-id=NXC-NO-TOT
subject=missing-total.txt
disposition=automatic-action/MDN-sent-automatically; processed
[log]
the message was PROCESSED
""",
        encoding="utf-8",
    )
    (sent_data / "no-details.txt").write_text(
        "CAB220 202600552                          EntregaFarm\nTOT00000000000\n",
        encoding="utf-8",
    )
    (sent / "no-details.hdr").write_text(
        """unique-id=NXC-NO-DET
subject=no-details.txt
disposition=automatic-action/MDN-sent-automatically; processed
""",
        encoding="utf-8",
    )
    (sent / "missing-edi.hdr").write_text(
        """unique-id=NXC-MISSING
subject=missing-edi.txt
disposition=automatic-action/MDN-sent-automatically; processed
""",
        encoding="utf-8",
    )
    (root / "received" / "data").mkdir(parents=True)
    (root / "received" / "data" / "not-processed.txt").write_text(
        "CAB220 202600552                          Farmacia Teste\nDET000001\nTOT00000000001\n",
        encoding="utf-8",
    )
    (received / "not-processed.hdr").write_text(
        """unique-id=NXC-PENDING
subject=ord.txt
[log]
message received.
""",
        encoding="utf-8",
    )
    (received / "error-log.hdr").write_text(
        """unique-id=NXC-ERROR
subject=ord-error.txt
[log]
cliente nao existe
""",
        encoding="utf-8",
    )

    written = export_header_report(root, tmp_path / "reports", run_id="exceptions")

    rows = {row["unique_id"]: row for row in json.loads(written[1].read_text(encoding="utf-8"))}
    assert rows["NXC-NO-TOT"]["exception_level"] == "warning"
    assert rows["NXC-NO-TOT"]["exception_reason"] == "edi_missing_total"
    assert rows["NXC-NO-DET"]["exception_level"] == "warning"
    assert rows["NXC-NO-DET"]["exception_reason"] == "edi_without_details"
    assert rows["NXC-MISSING"]["exception_level"] == "error"
    assert rows["NXC-MISSING"]["exception_reason"] == "edi_not_found"
    assert rows["NXC-PENDING"]["exception_level"] == "warning"
    assert rows["NXC-PENDING"]["exception_reason"] == "not_processed"
    assert rows["NXC-ERROR"]["exception_level"] == "error"
    assert rows["NXC-ERROR"]["exception_reason"] == "not_processed; edi_not_found; log_error"


def test_export_header_report_writes_exception_only_files(tmp_path: Path):
    root = tmp_path / "storage" / "cpip_1"
    sent = root / "sent" / "headers"
    sent_data = root / "sent" / "data"
    sent.mkdir(parents=True)
    sent_data.mkdir(parents=True)
    (sent_data / "ok.txt").write_text(
        "CAB220 202600552                          EntregaFarm\nDET000001\nTOT00000000001\n",
        encoding="utf-8",
    )
    (sent / "ok.hdr").write_text(
        """unique-id=NXC-OK
subject=ok.txt
disposition=automatic-action/MDN-sent-automatically; processed
""",
        encoding="utf-8",
    )
    (sent_data / "warning.txt").write_text(
        "CAB220 202600552                          EntregaFarm\nDET000001\n",
        encoding="utf-8",
    )
    (sent / "warning.hdr").write_text(
        """unique-id=NXC-WARNING
subject=warning.txt
disposition=automatic-action/MDN-sent-automatically; processed
""",
        encoding="utf-8",
    )

    written = export_header_report(root, tmp_path / "reports", run_id="split")

    exception_csv = tmp_path / "reports" / "split-exceptions.csv"
    exception_json = tmp_path / "reports" / "split-exceptions.json"
    assert exception_csv in written
    assert exception_json in written
    with exception_csv.open("r", newline="", encoding="utf-8") as handle:
        csv_rows = list(csv.DictReader(handle))
    json_rows = json.loads(exception_json.read_text(encoding="utf-8"))
    assert [row["unique_id"] for row in csv_rows] == ["NXC-WARNING"]
    assert [row["unique_id"] for row in json_rows] == ["NXC-WARNING"]


def test_export_header_report_writes_operational_xlsx(tmp_path: Path):
    root = tmp_path / "storage" / "cpip_1"
    sent = root / "sent" / "headers"
    sent_data = root / "sent" / "data"
    sent.mkdir(parents=True)
    sent_data.mkdir(parents=True)
    (sent_data / "ok.txt").write_text(
        "CAB220 202600552                          EntregaFarm\nDET000001\nTOT00000000001\n",
        encoding="utf-8",
    )
    (sent / "ok.hdr").write_text(
        """unique-id=NXC-OK
subject=ok.txt
disposition=automatic-action/MDN-sent-automatically; processed
""",
        encoding="utf-8",
    )
    (sent_data / "warning.txt").write_text(
        "CAB220 202600552                          EntregaFarm\nDET000001\n",
        encoding="utf-8",
    )
    (sent / "warning.hdr").write_text(
        """unique-id=NXC-WARNING
subject=warning.txt
disposition=automatic-action/MDN-sent-automatically; processed
""",
        encoding="utf-8",
    )

    written = export_header_report(root, tmp_path / "reports", run_id="excel")

    xlsx_path = tmp_path / "reports" / "excel.xlsx"
    assert xlsx_path in written
    workbook = load_workbook(xlsx_path)
    assert workbook.sheetnames == ["Todas", "Excepcoes", "Resumo"]
    assert workbook["Todas"].auto_filter.ref is not None
    assert workbook["Excepcoes"].max_row == 2
    assert workbook["Excepcoes"]["A2"].value == "sent"
    assert workbook["Resumo"]["A1"].value == "Estado"
    assert workbook["Resumo"]["B2"].value == 1
    assert workbook["Resumo"]["B3"].value == 1
    assert workbook["Resumo"]["B4"].value == 0
    assert workbook["Todas"]["L3"].fill.fill_type == "solid"
