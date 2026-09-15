from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from integration_app.generix.edi import GenerixEdiRecord, parse_edi_file
from integration_app.generix.headers import GenerixHeaderRecord, scan_header_dir


GENERIX_REPORT_COLUMNS = [
    "flow_type",
    "unique_id",
    "subject",
    "date",
    "from",
    "to",
    "message_id",
    "pipe_id",
    "receipt",
    "disposition",
    "processed",
    "exception_level",
    "exception_reason",
    "edi_path",
    "edi_origin_name",
    "edi_detail_count",
    "edi_has_total",
    "edi_gln_codes",
    "body_path",
    "header_path",
    "log_events",
]


def export_header_report(storage_root: Path, report_dir: Path, run_id: str) -> list[Path]:
    report_dir.mkdir(parents=True, exist_ok=True)
    records = _scan_storage_root(storage_root)
    rows = [_record_to_row(record, storage_root) for record in records]
    exception_rows = [row for row in rows if row["exception_level"] != "ok"]
    csv_path = report_dir / f"{run_id}.csv"
    json_path = report_dir / f"{run_id}.json"
    exception_csv_path = report_dir / f"{run_id}-exceptions.csv"
    exception_json_path = report_dir / f"{run_id}-exceptions.json"
    xlsx_path = report_dir / f"{run_id}.xlsx"
    _write_csv(csv_path, rows)
    _write_json(json_path, rows)
    _write_csv(exception_csv_path, exception_rows)
    _write_json(exception_json_path, exception_rows)
    _write_xlsx(xlsx_path, rows, exception_rows)
    return [csv_path, json_path, exception_csv_path, exception_json_path, xlsx_path]


def summarize_rows(rows: list[dict[str, object]]) -> dict[str, int]:
    return {
        "total": len(rows),
        "ok": _count_level(rows, "ok"),
        "warnings": _count_level(rows, "warning"),
        "errors": _count_level(rows, "error"),
    }


def _count_level(rows: list[dict[str, object]], level: str) -> int:
    return sum(1 for row in rows if row["exception_level"] == level)


def _scan_storage_root(storage_root: Path) -> list[GenerixHeaderRecord]:
    records: list[GenerixHeaderRecord] = []
    records.extend(scan_header_dir(storage_root / "received" / "headers", flow_type="received"))
    records.extend(scan_header_dir(storage_root / "sent" / "headers", flow_type="sent"))
    return records


def _record_to_row(record: GenerixHeaderRecord, storage_root: Path) -> dict[str, object]:
    edi_record = _parse_related_edi(record, storage_root)
    exception_level, exception_reason = _classify_exception(record, edi_record)
    return {
        "flow_type": record.flow_type,
        "unique_id": record.unique_id,
        "subject": record.subject,
        "date": record.date,
        "from": record.message_from,
        "to": record.message_to,
        "message_id": record.message_id,
        "pipe_id": record.pipe_id,
        "receipt": record.receipt,
        "disposition": record.disposition,
        "processed": record.processed,
        "exception_level": exception_level,
        "exception_reason": exception_reason,
        "edi_path": str(edi_record.path) if edi_record else None,
        "edi_origin_name": edi_record.origin_name if edi_record else None,
        "edi_detail_count": edi_record.detail_count if edi_record else None,
        "edi_has_total": edi_record.has_total if edi_record else None,
        "edi_gln_codes": edi_record.gln_codes if edi_record else [],
        "body_path": record.body_path,
        "header_path": str(record.header_path),
        "log_events": "\n".join(record.log_events),
    }


def _parse_related_edi(record: GenerixHeaderRecord, storage_root: Path) -> GenerixEdiRecord | None:
    path = _resolve_body_path(record, storage_root)
    if path is None:
        return None
    return parse_edi_file(path)


def _resolve_body_path(record: GenerixHeaderRecord, storage_root: Path) -> Path | None:
    if record.body_path:
        body_path = Path(record.body_path)
        if body_path.exists():
            return body_path
    fallback_path = storage_root / record.flow_type / "data" / f"{record.header_path.stem}.txt"
    if fallback_path.exists():
        return fallback_path
    return None


def _classify_exception(record: GenerixHeaderRecord, edi_record: GenerixEdiRecord | None) -> tuple[str, str]:
    reasons: list[str] = []
    if not record.processed:
        reasons.append("not_processed")
    if edi_record is None:
        reasons.append("edi_not_found")
    else:
        if edi_record.detail_count == 0:
            reasons.append("edi_without_details")
        if not edi_record.has_total:
            reasons.append("edi_missing_total")
    if _has_log_error(record):
        reasons.append("log_error")

    if not reasons:
        return "ok", ""
    if "edi_not_found" in reasons or "log_error" in reasons:
        return "error", "; ".join(reasons)
    return "warning", "; ".join(reasons)


def _has_log_error(record: GenerixHeaderRecord) -> bool:
    error_terms = ("error", "failed", "reject", "rejeit", "cliente nao existe", "cliente não existe")
    log_text = "\n".join(record.log_events).lower()
    return any(term in log_text for term in error_terms)


def _write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=GENERIX_REPORT_COLUMNS)
        writer.writeheader()
        writer.writerows(rows)


def _write_json(path: Path, rows: list[dict[str, object]]) -> None:
    path.write_text(json.dumps(rows, indent=2, ensure_ascii=False), encoding="utf-8")


def _write_xlsx(path: Path, rows: list[dict[str, object]], exception_rows: list[dict[str, object]]) -> None:
    workbook = Workbook()
    all_sheet = workbook.active
    all_sheet.title = "Todas"
    _write_rows_sheet(all_sheet, rows)
    exceptions_sheet = workbook.create_sheet("Excepcoes")
    _write_rows_sheet(exceptions_sheet, exception_rows)
    summary_sheet = workbook.create_sheet("Resumo")
    _write_summary_sheet(summary_sheet, summarize_rows(rows))
    workbook.save(path)


def _write_rows_sheet(sheet, rows: list[dict[str, object]]) -> None:
    sheet.append(GENERIX_REPORT_COLUMNS)
    for row in rows:
        sheet.append([_xlsx_value(row.get(column)) for column in GENERIX_REPORT_COLUMNS])
    _style_rows_sheet(sheet)


def _write_summary_sheet(sheet, summary: dict[str, int]) -> None:
    sheet.append(["Estado", "Total"])
    sheet.append(["OK", summary["ok"]])
    sheet.append(["Warnings", summary["warnings"]])
    sheet.append(["Errors", summary["errors"]])
    sheet.append(["Total", summary["total"]])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.column_dimensions["A"].width = 16
    sheet.column_dimensions["B"].width = 12


def _style_rows_sheet(sheet) -> None:
    if sheet.max_row == 0:
        return
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.auto_filter.ref = sheet.dimensions
    sheet.freeze_panes = "A2"
    for column_cells in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 60)
        sheet.column_dimensions[column_cells[0].column_letter].width = width
    exception_column = GENERIX_REPORT_COLUMNS.index("exception_level") + 1
    for row_number in range(2, sheet.max_row + 1):
        level = sheet.cell(row=row_number, column=exception_column).value
        fill = _exception_fill(level)
        if fill:
            for cell in sheet[row_number]:
                cell.fill = fill


def _exception_fill(level: object) -> PatternFill | None:
    if level == "warning":
        return PatternFill("solid", fgColor="FFF2CC")
    if level == "error":
        return PatternFill("solid", fgColor="F4CCCC")
    return None


def _xlsx_value(value: object) -> object:
    if isinstance(value, list):
        return "; ".join(str(item) for item in value)
    return value
